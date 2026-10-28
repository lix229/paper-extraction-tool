"""Load the paper list from the screening workbook and classify each source link."""
import re
from pathlib import Path
from urllib.parse import quote

import openpyxl

DOI_RE = re.compile(r'^10\.\d{4,9}/\S+$')
DOI_PREFIX_RE = re.compile(r'^(?:https?://(?:dx\.)?doi\.org/|doi:\s*)', re.I)
PUBMED_RE = re.compile(r'pubmed\.ncbi\.nlm\.nih\.gov/(\d+)')
RECORD_ID_RE = re.compile(r'^[A-Za-z0-9._-]{1,80}$')


def header_key(header):
    """Workbook headers look like 'review_aspect: What part of…'; keep the part before ':'."""
    return str(header).split(':', 1)[0].strip() if header is not None else ''


def clean(value):
    if value is None:
        return ''
    return str(value).strip()


def normalize_doi(text):
    doi = DOI_PREFIX_RE.sub('', text.strip())
    return doi if DOI_RE.match(doi) else None


def parse_source(raw):
    """Classify the workbook's DOI cell: a DOI string, a PubMed URL, another URL or blank."""
    raw = clean(raw)
    if not raw:
        return {'kind': 'missing', 'raw': '', 'url': None, 's2_id': None, 'label': 'No DOI'}
    match = PUBMED_RE.search(raw)
    if match:
        pmid = match.group(1)
        return {'kind': 'pmid', 'raw': raw, 'url': f'https://pubmed.ncbi.nlm.nih.gov/{pmid}/',
                's2_id': f'PMID:{pmid}', 'label': f'PubMed {pmid}'}
    doi = normalize_doi(raw)
    if doi:
        return {'kind': 'doi', 'raw': raw, 'doi': doi, 'url': 'https://doi.org/' + quote(doi, safe='/:;()._-'),
                's2_id': f'DOI:{doi}', 'label': f'DOI {doi}'}
    if raw.startswith(('http://', 'https://')):
        return {'kind': 'url', 'raw': raw, 'url': raw, 's2_id': None, 'label': 'Source link'}
    return {'kind': 'unknown', 'raw': raw, 'url': None, 's2_id': None, 'label': f'Unrecognised: {raw}'}


def load_papers(path, sheet, columns, screening):
    """Read one paper per row. `columns` maps id/title/authors/year/doi/abstract to header keys."""
    path = Path(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in workbook.sheetnames:
        raise ValueError(f'{path.name}: no sheet named {sheet!r}')
    rows = workbook[sheet].iter_rows(values_only=True)
    keys = [header_key(h) for h in next(rows)]
    absent = [c for c in [*columns.values(), *screening] if c not in keys]
    if absent:
        raise ValueError(f'{path.name}: missing columns {absent}')

    papers, seen = [], set()
    for row in rows:
        cells = dict(zip(keys, row))
        record_id = clean(cells.get(columns['id']))
        if not record_id:
            continue
        if not RECORD_ID_RE.match(record_id):
            raise ValueError(f'{path.name}: record id {record_id!r} must use letters, digits, ".", "_" or "-"')
        if record_id in seen:
            raise ValueError(f'{path.name}: duplicate record id {record_id}')
        seen.add(record_id)
        papers.append({
            'record_id': record_id,
            'title': clean(cells.get(columns['title'])),
            'authors': clean(cells.get(columns['authors'])),
            'year': clean(cells.get(columns['year'])),
            'abstract': clean(cells.get(columns['abstract'])),
            'source': parse_source(cells.get(columns['doi'])),
            'screening': {c: clean(cells.get(c)) for c in screening},
        })
    workbook.close()
    return papers
